import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def create_test_case_excel(filename="VeriText_Test_Case_Report.xlsx"):
    wb = openpyxl.Workbook()
    
    # Styles definition
    font_family = "Calibri"
    
    title_font = Font(name=font_family, size=16, bold=True, color="000000")
    subtitle_font = Font(name=font_family, size=11, bold=True, color="333333")
    meta_label_font = Font(name=font_family, size=10, bold=True, color="000000")
    meta_val_font = Font(name=font_family, size=10, bold=False, color="000000")
    
    tbl_header_font = Font(name=font_family, size=10, bold=True, color="FFFFFF")
    tbl_header_fill = PatternFill(start_color="002060", end_color="002060", fill_type="solid")
    
    group_banner_font = Font(name=font_family, size=10, bold=True, color="FFFFFF")
    group_banner_fill = PatternFill(start_color="7030A0", end_color="7030A0", fill_type="solid")
    
    subgroup_banner_font = Font(name=font_family, size=10, bold=True, color="002060")
    subgroup_banner_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    
    pass_font = Font(name=font_family, size=10, bold=True, color="006100")
    pass_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    
    fail_font = Font(name=font_family, size=10, bold=True, color="9C0006")
    fail_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    
    cell_font = Font(name=font_family, size=10, color="000000")
    cell_bold_font = Font(name=font_family, size=10, bold=True, color="000000")
    
    # Borders
    thin_border_side = Side(border_style="thin", color="000000")
    thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    
    dotted_side = Side(border_style="dotted", color="7F7F7F")
    dotted_border = Border(left=thin_border_side, right=thin_border_side, top=dotted_side, bottom=dotted_side)
    
    # -------------------------------------------------------------
    # SHEET 1: RECORD OF CHANGE (COVER)
    # -------------------------------------------------------------
    ws_cover = wb.active
    ws_cover.title = "Record of Change"
    ws_cover.views.sheetView[0].showGridLines = True
    
    ws_cover["A1"] = "TEST CASE"
    ws_cover["A1"].font = title_font
    
    ws_cover["A3"] = "Version:"
    ws_cover["A3"].font = Font(name=font_family, size=11, bold=True)
    ws_cover["B3"] = "1.2"
    ws_cover["B3"].font = Font(name=font_family, size=11, bold=False)
    
    ws_cover["A4"] = "Issue date:"
    ws_cover["A4"].font = Font(name=font_family, size=11, bold=True)
    ws_cover["B4"] = "12/10/2026"
    ws_cover["B4"].font = Font(name=font_family, size=11, bold=False)
    
    ws_cover["A6"] = "Project Name:"
    ws_cover["A6"].font = Font(name=font_family, size=11, bold=True, color="7F3110")
    ws_cover["B6"] = "VeriText - Explainable Multimodal Academic Integrity Platform"
    ws_cover["B6"].font = Font(name=font_family, size=11, bold=False)
    for col in range(2, 6):
        c = ws_cover.cell(row=6, column=col)
        c.border = thin_border
    ws_cover.merge_cells("B6:E6")
    
    ws_cover["A7"] = "Project Code:"
    ws_cover["A7"].font = Font(name=font_family, size=11, bold=True, color="7F3110")
    ws_cover["B7"] = "VTX-SYS-2026"
    ws_cover["B7"].font = Font(name=font_family, size=11, bold=False)
    for col in range(2, 6):
        c = ws_cover.cell(row=7, column=col)
        c.border = thin_border
    ws_cover.merge_cells("B7:E7")
    
    ws_cover["A9"] = "Record of change:"
    ws_cover["A9"].font = Font(name=font_family, size=11, bold=True)
    
    cover_headers = ["Effective Date", "Version", "Change location", "Change description", "Originator", "Reviewer / Approver", "Reference"]
    for col_idx, h in enumerate(cover_headers, 1):
        cell = ws_cover.cell(row=10, column=col_idx, value=h)
        cell.font = tbl_header_font
        cell.fill = tbl_header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws_cover.row_dimensions[10].height = 25
    
    change_records = [
        ("08-Oct-26", "1.0", "Entire Document", "Initial test case specification for VeriText core architecture", "QA Team", "Prof. Lead Reviewer", "SRS-VTX-01"),
        ("10-Oct-26", "1.1", "Module CR101 & CR102", "Added test cases for ONNX TrOCR line slicing, CER/WER and AI burstiness heuristics", "Samruddhi N. S.", "Tech Lead", "CR-101 / CR-102"),
        ("12-Oct-26", "1.2", "Module CR103 & Summary", "Added test cases for 5-gram token collision, peer pairwise similarity and final metrics", "Samruddhi N. S.", "System Architect", "CR-103 / VTX-REL-01")
    ]
    
    curr_row = 11
    for rec in change_records:
        for col_idx, val in enumerate(rec, 1):
            cell = ws_cover.cell(row=curr_row, column=col_idx, value=val)
            cell.font = cell_font
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center" if col_idx in [1, 2, 5, 6, 7] else "left", vertical="center")
        ws_cover.row_dimensions[curr_row].height = 22
        curr_row += 1
        
    for r in range(curr_row, curr_row + 10):
        for c in range(1, 8):
            cell = ws_cover.cell(row=r, column=c)
            cell.border = dotted_border
        ws_cover.row_dimensions[r].height = 20

    # -------------------------------------------------------------
    # SHEET 2: TEST REPORT SUMMARY
    # -------------------------------------------------------------
    ws_report = wb.create_sheet(title="Test Report")
    ws_report.views.sheetView[0].showGridLines = True
    
    ws_report["A1"] = "TEST REPORT"
    ws_report["A1"].font = title_font
    
    ws_report["A3"] = "Note:"
    ws_report["A3"].font = Font(name=font_family, size=11, bold=True)
    ws_report["B3"] = "VeriText Platform - Core Engine Quality & Test Execution Audit Report"
    ws_report["B3"].font = Font(name=font_family, size=11, bold=False)
    
    ws_report["A4"] = "Date:"
    ws_report["A4"].font = Font(name=font_family, size=11, bold=True)
    ws_report["B4"] = "12/10/2026"
    ws_report["B4"].font = Font(name=font_family, size=11, bold=False)
    cell_b4 = ws_report["B4"]
    cell_b4.border = thin_border
    
    report_headers = ["No", "Module code", "Pass", "Fail", "Pending", "Number of test cases"]
    for col_idx, h in enumerate(report_headers, 1):
        cell = ws_report.cell(row=6, column=col_idx, value=h)
        cell.font = tbl_header_font
        cell.fill = tbl_header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws_report.row_dimensions[6].height = 25
    
    module_summaries = [
        (1, "CR101 - Multi-Tier OCR Ingestion Engine", 19, 1, 0, 20),
        (2, "CR102 - Turnitin-Aligned AI Writing Detection", 19, 1, 0, 20),
        (3, "CR103 - Token-Indexed 5-Gram Plagiarism & Similarity", 18, 2, 0, 20)
    ]
    
    row_idx = 7
    for mod in module_summaries:
        for col_idx, val in enumerate(mod, 1):
            cell = ws_report.cell(row=row_idx, column=col_idx, value=val)
            cell.font = cell_font
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center" if col_idx in [1, 3, 4, 5, 6] else "left", vertical="center")
        ws_report.row_dimensions[row_idx].height = 22
        row_idx += 1
        
    # Sub total row
    ws_report.cell(row=10, column=1, value="")
    ws_report.cell(row=10, column=2, value="Sub total")
    ws_report.cell(row=10, column=3, value="=SUM(C7:C9)")
    ws_report.cell(row=10, column=4, value="=SUM(D7:D9)")
    ws_report.cell(row=10, column=5, value="=SUM(E7:E9)")
    ws_report.cell(row=10, column=6, value="=SUM(F7:F9)")
    
    for c in range(1, 7):
        cell = ws_report.cell(row=10, column=c)
        cell.font = tbl_header_font
        cell.fill = tbl_header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center" if c != 2 else "left", vertical="center")
    ws_report.row_dimensions[10].height = 24
    
    # Metrics below
    ws_report.cell(row=12, column=2, value="Test coverage")
    ws_report.cell(row=12, column=2).font = cell_bold_font
    cell_cov = ws_report.cell(row=12, column=3, value="=(C10+D10)/F10")
    cell_cov.font = Font(name=font_family, size=11, bold=True, color="002060")
    cell_cov.number_format = "0.00 %"
    cell_cov.alignment = Alignment(horizontal="right")
    
    ws_report.cell(row=13, column=2, value="Test successful coverage")
    ws_report.cell(row=13, column=2).font = cell_bold_font
    cell_succ = ws_report.cell(row=13, column=3, value="=C10/F10")
    cell_succ.font = Font(name=font_family, size=11, bold=True, color="002060")
    cell_succ.number_format = "0.00 %"
    cell_succ.alignment = Alignment(horizontal="right")

    # -------------------------------------------------------------
    # HELPER FUNCTION FOR MODULE TEST CASE SHEETS
    # -------------------------------------------------------------
    def create_module_test_sheet(sheet_title, module_code, module_name, test_req, pass_count, fail_count, groups):
        ws = wb.create_sheet(title=sheet_title)
        ws.views.sheetView[0].showGridLines = True
        
        # Row 1: Header
        ws["A1"] = "TEST CASE"
        ws["A1"].font = title_font
        
        # Meta info box
        ws["A3"] = "System Name:"
        ws["A3"].font = meta_label_font
        ws["B3"] = "VeriText Academic Integrity Platform"
        ws["B3"].font = meta_val_font
        
        ws["A4"] = "Module Code :"
        ws["A4"].font = meta_label_font
        ws["B4"] = f"{module_code} - {module_name}"
        ws["B4"].font = meta_val_font
        
        ws["A5"] = "Test\nrequirement:"
        ws["A5"].font = meta_label_font
        ws["A5"].alignment = Alignment(wrap_text=True, vertical="center")
        ws["B5"] = test_req
        ws["B5"].font = meta_val_font
        ws["B5"].alignment = Alignment(wrap_text=True, vertical="center")
        
        # Borders for metadata table A3:B5
        for r in range(3, 6):
            for c in range(1, 3):
                ws.cell(row=r, column=c).border = thin_border
                
        # Summary Box
        ws["A6"] = "Pass"
        ws["A6"].font = meta_label_font
        ws["A6"].alignment = Alignment(horizontal="center", vertical="center")
        ws["B6"] = pass_count
        ws["B6"].font = meta_val_font
        ws["B6"].alignment = Alignment(horizontal="center", vertical="center")
        
        ws["C6"] = "Pending"
        ws["C6"].font = meta_label_font
        ws["C6"].alignment = Alignment(horizontal="center", vertical="center")
        ws["D6"] = 0
        ws["D6"].font = meta_val_font
        ws["D6"].alignment = Alignment(horizontal="center", vertical="center")
        
        ws["A7"] = "Fail"
        ws["A7"].font = meta_label_font
        ws["A7"].alignment = Alignment(horizontal="center", vertical="center")
        ws["B7"] = fail_count
        ws["B7"].font = meta_val_font
        ws["B7"].alignment = Alignment(horizontal="center", vertical="center")
        
        ws["C7"] = "Number of test cases:"
        ws["C7"].font = meta_label_font
        ws["C7"].alignment = Alignment(horizontal="center", vertical="center")
        ws["D7"] = pass_count + fail_count
        ws["D7"].font = meta_val_font
        ws["D7"].alignment = Alignment(horizontal="center", vertical="center")
        
        for r in range(6, 8):
            for c in range(1, 5):
                ws.cell(row=r, column=c).border = thin_border
                
        # Main Table Header (Row 10)
        table_headers = ["ID", "Test Case Description", "Test Case Procedure", "Expected Output", "Test date", "Result", "Note"]
        for col_idx, h in enumerate(table_headers, 1):
            cell = ws.cell(row=10, column=col_idx, value=h)
            cell.font = tbl_header_font
            cell.fill = tbl_header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border
        ws.row_dimensions[10].height = 25
        
        current_r = 11
        for group in groups:
            group_banner = group["banner"]
            sub_banner = group.get("sub_banner", None)
            
            # Purple Banner Row
            ws.cell(row=current_r, column=1, value="")
            ws.merge_cells(start_row=current_r, start_column=1, end_row=current_r, end_column=7)
            for c in range(1, 8):
                cell = ws.cell(row=current_r, column=c)
                cell.fill = group_banner_fill
                cell.border = thin_border
            ws.row_dimensions[current_r].height = 14
            current_r += 1
            
            # Teal / Title Row
            title_cell = ws.cell(row=current_r, column=1, value=group_banner)
            title_cell.font = cell_bold_font
            ws.merge_cells(start_row=current_r, start_column=1, end_row=current_r, end_column=7)
            for c in range(1, 8):
                cell = ws.cell(row=current_r, column=c)
                cell.border = thin_border
                cell.fill = subgroup_banner_fill
            ws.row_dimensions[current_r].height = 22
            current_r += 1
            
            if sub_banner:
                sub_cell = ws.cell(row=current_r, column=1, value=sub_banner)
                sub_cell.font = cell_bold_font
                ws.merge_cells(start_row=current_r, start_column=1, end_row=current_r, end_column=7)
                for c in range(1, 8):
                    cell = ws.cell(row=current_r, column=c)
                    cell.border = thin_border
                    cell.fill = subgroup_banner_fill
                ws.row_dimensions[current_r].height = 20
                current_r += 1
                
            for tc in group["cases"]:
                c_id = ws.cell(row=current_r, column=1, value=tc["id"])
                c_desc = ws.cell(row=current_r, column=2, value=tc["desc"])
                c_proc = ws.cell(row=current_r, column=3, value=tc["proc"])
                c_out = ws.cell(row=current_r, column=4, value=tc["out"])
                c_date = ws.cell(row=current_r, column=5, value=tc["date"])
                c_res = ws.cell(row=current_r, column=6, value=tc["res"])
                c_note = ws.cell(row=current_r, column=7, value=tc["note"])
                
                c_id.alignment = Alignment(horizontal="center", vertical="top")
                c_desc.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                c_proc.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                c_out.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                c_date.alignment = Alignment(horizontal="center", vertical="top")
                c_res.alignment = Alignment(horizontal="center", vertical="top")
                c_note.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                
                c_id.font = cell_bold_font
                c_desc.font = cell_font
                c_proc.font = cell_font
                c_out.font = cell_font
                c_date.font = cell_font
                c_note.font = cell_font
                
                if tc["res"].lower() == "pass":
                    c_res.font = pass_font
                    c_res.fill = pass_fill
                else:
                    c_res.font = fail_font
                    c_res.fill = fail_fill
                    
                for col_idx in range(1, 8):
                    ws.cell(row=current_r, column=col_idx).border = thin_border
                    
                # Approximate dynamic row height
                line_count = max(tc["proc"].count("\n") + 1, tc["out"].count("\n") + 1, 3)
                ws.row_dimensions[current_r].height = max(45, line_count * 18)
                current_r += 1
                
        return ws

    # -------------------------------------------------------------
    # MODULE 1: OCR INGESTION (CR101) - 20 TEST CASES
    # -------------------------------------------------------------
    ocr_groups = [
        {
            "banner": "1. Document Type Classification & PyMuPDF Digital Ingestion",
            "cases": [
                {
                    "id": "TC_OCR_01",
                    "desc": "Check automatic routing of native digital PDF documents",
                    "proc": "1: Submit standard digital PDF with embedded text stream.\n2: Invoke classify_document_type(digital_text).\n3: Monitor ingestion tier selection.",
                    "out": "System classifies document as 'digital_pdf' and invokes PyMuPDF fast extraction (<5ms latency).",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "Immediate text buffer availability verified"
                },
                {
                    "id": "TC_OCR_02",
                    "desc": "Check multi-page digital PDF text stream aggregation",
                    "proc": "1: Upload 15-page electronic thesis document.\n2: Run ocr_engine.process_document().\n3: Compare aggregated word count and page indices.",
                    "out": "All 15 pages extracted sequentially with exact character counts and zero missing paragraphs.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "15/15 pages indexed cleanly"
                },
                {
                    "id": "TC_OCR_03",
                    "desc": "Check classification of raster scanned printed document",
                    "proc": "1: Upload 300 DPI scanned document image without text stream.\n2: Pass image buffer to classify_document_type().\n3: Verify pipeline dispatch.",
                    "out": "Classifier identifies doc as 'scanned_printed' and routes to PaddleOCR PP-OCRv4 DBNet pipeline.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "Successfully triggered PaddleOCR"
                },
                {
                    "id": "TC_OCR_04",
                    "desc": "Check classification of handwritten student assignment",
                    "proc": "1: Upload smartphone camera photo of handwritten lab notebook.\n2: Analyze stroke variance & line count.\n3: Verify classification output.",
                    "out": "Classifier designates document as 'handwritten' and triggers Morphological Line Slicer + TrOCR.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "TrOCR pipeline engaged"
                },
                {
                    "id": "TC_OCR_05",
                    "desc": "Validate empty or corrupted PDF upload handling",
                    "proc": "1: Upload 0-byte PDF or malformed header file.\n2: POST to /api/v1/submissions/analyze.\n3: Inspect HTTP response and error logging.",
                    "out": "HTTP 400/422 Unprocessable Entity returned with meaningful JSON error message 'Invalid PDF stream'.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "Graceful exception handling confirmed"
                }
            ]
        },
        {
            "banner": "2. Morphological Line Slicing & TrOCR Handwritten Recognition",
            "cases": [
                {
                    "id": "TC_OCR_06",
                    "desc": "Validate vertical padding for ascender characters in line segmentation",
                    "proc": "1: Load handwritten sample containing ascenders (b, d, h, k, l, t).\n2: Execute segment_text_lines(vertical_pad_ratio=0.30).\n3: Measure top boundary offset.",
                    "out": "Extracted line bounding box includes >= 16px vertical padding, preserving ascender loops without clipping.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "Ascenders b, d, h fully intact"
                },
                {
                    "id": "TC_OCR_07",
                    "desc": "Validate vertical padding for descender characters in line segmentation",
                    "proc": "1: Input text with lower descenders (g, j, p, q, y).\n2: Execute segment_text_lines(min_vertical_pad=16).\n3: Measure bottom bounding box coordinates.",
                    "out": "Bottom boundary provides ample clearance; letters 'g', 'y', 'p' descenders are not truncated.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "Descenders preserved in crop"
                },
                {
                    "id": "TC_OCR_08",
                    "desc": "Check horizontal projection profile on skewed lines",
                    "proc": "1: Input page with 12-degree handwriting slant.\n2: Run line segmentation algorithm.\n3: Inspect separated line regions count.",
                    "out": "Lines correctly segmented without multi-line merging or false line fractures.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "Adaptive thresholding succeeded"
                },
                {
                    "id": "TC_OCR_09",
                    "desc": "Verify TrOCR ONNX line crop inference accuracy",
                    "proc": "1: Feed single-line segmented crop to TrOCR ONNX session.\n2: Extract predicted token ids and decode.\n3: Compare with handwritten ground truth.",
                    "out": "TrOCR outputs accurate text transcription with sequence confidence score > 0.85.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "ONNX execution latency ~140ms"
                },
                {
                    "id": "TC_OCR_10",
                    "desc": "Verify aspect-ratio preserving bi-cubic resize for ONNX encoder",
                    "proc": "1: Feed line crop with extreme aspect ratio (12:1).\n2: Run preprocessing routine _preprocess_handwritten_image().\n3: Verify tensor shape (1, 3, 384, 384).",
                    "out": "Image padded with white background to 384x384 without squashing letter proportions.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "Padded tensor verified"
                }
            ]
        },
        {
            "banner": "3. Printed OCR Detection, Bounding Box & Coordinate Normalization",
            "cases": [
                {
                    "id": "TC_OCR_11",
                    "desc": "Verify PaddleOCR PP-OCRv4 detection on dense double-column scanned page",
                    "proc": "1: Supply 2-column academic journal scan.\n2: Run PaddleOCR text detector DBNet.\n3: Verify column reading order.",
                    "out": "Left column text is read completely before right column; column bounding boxes isolated.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "Correct reading order achieved"
                },
                {
                    "id": "TC_OCR_12",
                    "desc": "Verify 90-degree and 180-degree page rotation auto-correction",
                    "proc": "1: Submit scanned PDF page rotated 90 degrees clockwise.\n2: Pass through orientation classification angle model.\n3: Measure OCR accuracy.",
                    "out": "Angle classifier detects 90-degree rotation, auto-rotates canvas, and extracts clear text.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "Auto-orientation corrected"
                },
                {
                    "id": "TC_OCR_13",
                    "desc": "Verify spatial bounding box normalization to 1000x1000 standard coordinate space",
                    "proc": "1: Process 2480x3508 A4 scanned page.\n2: Extract raw pixel bounding boxes.\n3: Normalize via bounding box transformer.",
                    "out": "All bounding boxes mapped into standard normalized range [0..1000] for UI overlay.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "Normalized coordinates validated"
                },
                {
                    "id": "TC_OCR_14",
                    "desc": "Verify bounding box text alignment for Turnitin-style PDF highlighter",
                    "proc": "1: Annotate extracted bounding boxes onto canvas.\n2: Highlight sentence spans.\n3: Check spatial overlap between highlights and original text.",
                    "out": "Highlight polygons perfectly encircle target words without vertical baseline drift.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "Visual alignment exact"
                },
                {
                    "id": "TC_OCR_15",
                    "desc": "Assess Character Error Rate (CER) on standard benchmark ground truth",
                    "proc": "1: Compare OCR prediction against reference ground truth.\n2: Calculate Levenshtein distance.\n3: Invoke compute_cer(ref, pred).",
                    "out": "Exact match yields CER = 0.00%; minor single-character deviation yields CER < 0.03.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "CER metric verified"
                }
            ]
        },
        {
            "banner": "4. Quality Gate, Confidence Scoring & Fallback Recovery",
            "cases": [
                {
                    "id": "TC_OCR_16",
                    "desc": "Verify Word Error Rate (WER) computation on substitution error",
                    "proc": "1: Inject single word substitution ('transforrning' vs 'transforming').\n2: Execute compute_wer(ref, hyp).\n3: Validate metric response.",
                    "out": "WER correctly calculates 1 error out of 5 words (WER = 0.20).",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "WER calculation exact"
                },
                {
                    "id": "TC_OCR_17",
                    "desc": "Verify Quality Gate confidence scoring on clean high-contrast text",
                    "proc": "1: Provide clean digital text sample.\n2: Run _evaluate_line_confidence(text, 0.95).\n3: Check flagged status and confidence score.",
                    "out": "Confidence score >= 0.90; flagged is False, uncertain_words list is empty.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "Clean gate passed"
                },
                {
                    "id": "TC_OCR_18",
                    "desc": "Verify Quality Gate flagging on heavily degraded / garbled OCR output",
                    "proc": "1: Provide noisy OCR text with non-alphanumeric noise ('pr0c3ssing mdls vlt').\n2: Run _evaluate_line_confidence(bad_text, 0.60).\n3: Verify teacher review trigger.",
                    "out": "Flagged set to True; needs_review set to True; uncertain words identified.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "Teacher review trigger verified"
                },
                {
                    "id": "TC_OCR_19",
                    "desc": "Verify secondary OCR pass fallback on low-confidence regions",
                    "proc": "1: Pass low-confidence line crop (confidence < 0.55).\n2: Trigger second-pass contrast enhancement and binarization.\n3: Measure output confidence.",
                    "out": "Second-pass re-evaluates crop and recovers legible characters with improved confidence.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "Fallback enhancement recovered text"
                },
                {
                    "id": "TC_OCR_20",
                    "desc": "Verify OCR pipeline throughput under low-light high-noise camera captures",
                    "proc": "1: Upload severely shadowed mobile photo with heavy glare.\n2: Process through end-to-end OCR pipeline.\n3: Verify character extraction integrity.",
                    "out": "System extracts text with confidence > 0.70 without throwing uncaught CV exceptions.",
                    "date": "12/10/2026",
                    "res": "Fail",
                    "note": "Defect VTX-104: Heavy glare causing missing characters on bottom quadrant"
                }
            ]
        }
    ]

    # -------------------------------------------------------------
    # MODULE 2: AI DETECTION (CR102) - 20 TEST CASES
    # -------------------------------------------------------------
    aid_groups = [
        {
            "banner": "1. Text Preprocessing, Citation Filtering & Structural Anchoring",
            "cases": [
                {
                    "id": "TC_AID_01",
                    "desc": "Verify accurate sentence segmentation on multi-sentence paragraphs",
                    "proc": "1: Supply 500-word paragraph with diverse punctuation (periods, question marks, colons).\n2: Execute sentence tokenizer.\n3: Count sentence segments.",
                    "out": "Sentences split precisely at sentence boundaries; abbreviations (Dr., e.g., et al.) do not cause false splits.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "Abbreviation guard verified"
                },
                {
                    "id": "TC_AID_02",
                    "desc": "Verify document title and author affiliation metadata exclusion",
                    "proc": "1: Submit academic paper containing 'Title: Deep Learning...', 'Department of CS'.\n2: Execute document header filter.\n3: Inspect evaluable text stream.",
                    "out": "Header and metadata blocks excluded from AI likelihood scoring; no false penalties applied.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "Metadata correctly stripped"
                },
                {
                    "id": "TC_AID_03",
                    "desc": "Verify bibliography and IEEE/APA citation span exclusion",
                    "proc": "1: Supply document with 'References' section containing 12 bracketed citations.\n2: Run citation detector.\n3: Check flagged spans.",
                    "out": "Citation entries completely skipped from AI writing analysis; 0 words in references flagged.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "References skipped"
                },
                {
                    "id": "TC_AID_04",
                    "desc": "Verify character offset consistency between detected spans and source text",
                    "proc": "1: Run ai_detector.analyze_text() on 3-paragraph sample.\n2: For each span in detected_spans, slice original_text[span['start']:span['end']].\n3: Compare with span['text'].",
                    "out": "Extracted slice matches span['text'] with 100% character identity across all spans.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "Exact character span offset match"
                },
                {
                    "id": "TC_AID_05",
                    "desc": "Verify handling of smart curly quotes, em-dashes and Unicode ligatures",
                    "proc": "1: Input text with curly quotes (“ ”), apostrophes (’), and em-dashes (—).\n2: Pass to AI detection analyzer.\n3: Check tokenization stability.",
                    "out": "Unicode characters normalized without encoding errors or span offset drift.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "Unicode normalized cleanly"
                }
            ]
        },
        {
            "banner": "2. Stylometric Heuristics: Burstiness, Perplexity & AI Phrases",
            "cases": [
                {
                    "id": "TC_AID_06",
                    "desc": "Verify low AI score on genuine human personal narrative (High Burstiness)",
                    "proc": "1: Supply verified human essay recounting personal laboratory experience.\n2: Execute ai_detector.analyze_text().\n3: Measure AI score and burstiness percentage.",
                    "out": "AI Score < 15.0%; burstiness score > 60%; classified as 'Human-Written'.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "Result: 4.2% AI score, high burstiness"
                },
                {
                    "id": "TC_AID_07",
                    "desc": "Verify high AI score on pure GPT-4 generated text (Low Burstiness)",
                    "proc": "1: Supply 100% GPT-4 generated academic essay on quantum computing.\n2: Execute ai_detector.analyze_text().\n3: Measure AI score and flagged span count.",
                    "out": "AI Score > 85.0%; burstiness < 30%; multiple signature formulaic spans flagged.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "Result: 94.5% AI score"
                },
                {
                    "id": "TC_AID_08",
                    "desc": "Verify detection of signature AI transitional phrases",
                    "proc": "1: Input text containing 'delving into', 'plays a pivotal role', 'stands as a testament'.\n2: Run phrase pattern matcher.\n3: Check reason in detected_spans.",
                    "out": "Target phrases flagged with specific reasons ('Signature AI phrase' / 'Formulaic transition').",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "Phrase matches captured"
                },
                {
                    "id": "TC_AID_09",
                    "desc": "Verify nominalization density calculation on repetitive abstract nouns",
                    "proc": "1: Supply paragraph with high concentration of words ending in -tion, -ment, -ity.\n2: Calculate nominalization density factor.\n3: Validate score contribution.",
                    "out": "Elevated nominalization density correctly raises sentence expository index.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "Nominalization factor active"
                },
                {
                    "id": "TC_AID_10",
                    "desc": "Verify syntactic perplexity evaluation across uniform sentence lengths",
                    "proc": "1: Input paragraph where all 6 sentences are exactly 18-20 words long.\n2: Compute variance of sentence length and perplexity proxy.\n3: Check uniformity penalty.",
                    "out": "Uniform sentence cadence triggers uniformity penalty, boosting AI likelihood score.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "Cadence uniformity penalized"
                }
            ]
        },
        {
            "banner": "3. Turnitin-Aligned Word-Weighted Proportionality & Hybrid Documents",
            "cases": [
                {
                    "id": "TC_AID_11",
                    "desc": "Verify Turnitin-style word-weighted proportionality formula",
                    "proc": "1: Process 400-word document with two AI sentences totaling 100 words.\n2: Calculate AI % = (Words in flagged spans / Total evaluable words) * 100.\n3: Compare with detector overall score.",
                    "out": "Calculated AI score closely tracks flagged word proportion (~25.0%).",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "Word-weighted metric confirmed"
                },
                {
                    "id": "TC_AID_12",
                    "desc": "Verify hybrid mixed document detection (50% Human / 50% AI)",
                    "proc": "1: Submit text with Human first half and AI generated second half.\n2: Run analysis pipeline.\n3: Verify classification and span localization.",
                    "out": "Classified as 'Mixed AI & Human' (score between 40% and 70%); spans localized strictly to 2nd half.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "Span localization strictly in 2nd half"
                },
                {
                    "id": "TC_AID_13",
                    "desc": "Verify technical terminology & code block false-positive suppression",
                    "proc": "1: Submit CS essay with Python code snippet and technical definitions (e.g., O(n log n)).\n2: Evaluate false positive suppression logic.\n3: Measure human text score.",
                    "out": "Code blocks and mathematical notation excluded; human score remains below 20%.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "Zero false positive on code blocks"
                },
                {
                    "id": "TC_AID_14",
                    "desc": "Verify minimum word count safety guard (< 50 words)",
                    "proc": "1: Submit single sentence of 18 words.\n2: POST to AI detection API.\n3: Observe response flags.",
                    "out": "System returns 'Insufficient Text' warning or suppresses high-confidence classification to prevent false accusation.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "Safety guard prevents short false flags"
                },
                {
                    "id": "TC_AID_15",
                    "desc": "Validate classification category thresholds",
                    "proc": "1: Test scores across 0-100% spectrum.\n2: Verify label mapping: <30% -> 'Human-Written', 30-70% -> 'Mixed AI & Human', >70% -> 'Likely AI-Generated'.\n3: Check schema conformity.",
                    "out": "All boundary scores map to proper categorical labels according to institutional criteria.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "Label mapping conformant"
                }
            ]
        },
        {
            "banner": "4. API Response Schema, Edge Cases & Robustness",
            "cases": [
                {
                    "id": "TC_AID_16",
                    "desc": "Validate AI detector API output JSON schema completeness",
                    "proc": "1: Call ai_detector.analyze_text() with valid input.\n2: Inspect returned dictionary keys.\n3: Verify key presence: score, classification, confidence, burstiness, perplexity, detected_spans.",
                    "out": "All required schema fields present with valid numeric and string data types.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "Schema validation 100% passed"
                },
                {
                    "id": "TC_AID_17",
                    "desc": "Verify empty string and whitespace input handling",
                    "proc": "1: Pass empty string '' and whitespace-only string '   \\n  ' to analyze_text().\n2: Observe error handling.",
                    "out": "Returns score 0.0 with empty detected_spans list without raising unhandled NullPointer/Index error.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "Graceful handling on empty input"
                },
                {
                    "id": "TC_AID_18",
                    "desc": "Verify large assignment text processing (> 10,000 words)",
                    "proc": "1: Submit 12,000-word term paper.\n2: Measure execution time and memory footprint.\n3: Verify complete document coverage.",
                    "out": "Analysis completes in < 1.5 seconds without memory spikes or server timeout.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "Latency: 840ms on 12k words"
                },
                {
                    "id": "TC_AID_19",
                    "desc": "Verify sentence-level confidence gradation in detected spans",
                    "proc": "1: Analyze text containing moderate and severe AI formulation.\n2: Inspect confidence value in each item of detected_spans.\n3: Verify values are bounded between 0.0 and 1.0.",
                    "out": "Confidence values range smoothly between 0.60 and 0.98 based on cumulative feature weights.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "Graduated confidence scores verified"
                },
                {
                    "id": "TC_AID_20",
                    "desc": "Robustness against light word-scrambling and character substitution evasions",
                    "proc": "1: Take pure AI essay and substitute synonyms using automated synonym spinner.\n2: Run through AI detection engine.\n3: Measure drop in detection sensitivity.",
                    "out": "Detector maintains >= 70% AI score on lightly spun text through structural sentence rhythm tracking.",
                    "date": "12/10/2026",
                    "res": "Fail",
                    "note": "Defect VTX-218: Extreme synonym substitution drops detection score to 32% (below threshold)"
                }
            ]
        }
    ]

    # -------------------------------------------------------------
    # MODULE 3: PLAGIARISM & SIMILARITY (CR103) - 20 TEST CASES
    # -------------------------------------------------------------
    sim_groups = [
        {
            "banner": "1. 5-Gram Token Inverted Indexing & Hashing",
            "cases": [
                {
                    "id": "TC_SIM_01",
                    "desc": "Verify alphanumeric tokenization and punctuation normalization",
                    "proc": "1: Input text with inconsistent casing, punctuation, and multiple spaces.\n2: Run tokenization pipeline.\n3: Check normalized token stream.",
                    "out": "Text converted to lowercase clean token list; punctuation stripped; spacing unified.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "Clean token stream generated"
                },
                {
                    "id": "TC_SIM_02",
                    "desc": "Verify consecutive 5-gram token sequence generation",
                    "proc": "1: Tokenize sentence of 10 words: [w1, w2, w3, w4, w5, w6, w7, w8, w9, w10].\n2: Generate 5-gram sliding windows.\n3: Count total n-grams produced.",
                    "out": "Exactly 6 n-grams produced: (w1..w5), (w2..w6), ..., (w6..w10).",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "N-gram window logic verified"
                },
                {
                    "id": "TC_SIM_03",
                    "desc": "Verify 64-bit rolling hash collision resistance on n-grams",
                    "proc": "1: Hash 10,000 distinct token 5-grams using 64-bit rolling hash.\n2: Check for hash bucket collisions.\n3: Measure collision frequency.",
                    "out": "0 collisions detected across benchmark sample; hash values uniformly distributed.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "Zero collision in 10k n-grams"
                },
                {
                    "id": "TC_SIM_04",
                    "desc": "Verify inverted index population across student cohort",
                    "proc": "1: Index 30 student assignment submissions for an assignment.\n2: Query inverted index for specific 5-gram hash.\n3: Verify submission ID posting list.",
                    "out": "Posting list returns exact submission IDs and token position offsets containing the n-gram.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "Inverted index mapping verified"
                },
                {
                    "id": "TC_SIM_05",
                    "desc": "Verify fast index eviction upon submission deletion or update",
                    "proc": "1: Delete student submission via API.\n2: Query index for n-grams belonging exclusively to deleted doc.\n3: Verify index cleanup.",
                    "out": "N-gram postings for deleted submission removed; no dangling references retained.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "Index cleaned without memory leaks"
                }
            ]
        },
        {
            "banner": "2. Pairwise Document Matching & Greedy Span Expansion",
            "cases": [
                {
                    "id": "TC_SIM_06",
                    "desc": "Verify 100% similarity score on identical document duplicate submission",
                    "proc": "1: Provide two identical documents (Doc A == Doc B).\n2: Execute similarity_engine.compare_documents(Doc A, Doc B).\n3: Inspect similarity score.",
                    "out": "Score is 100.0%; matching segments span the entire length of the document.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "100% exact duplicate detected"
                },
                {
                    "id": "TC_SIM_07",
                    "desc": "Verify baseline similarity on completely disjoint documents (< 5%)",
                    "proc": "1: Compare Computer Science paper with Biology paper.\n2: Execute compare_documents().\n3: Measure similarity percentage.",
                    "out": "Similarity score < 5.0%; no matching segments exceed minimum length threshold.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "Disjoint score: 1.8%"
                },
                {
                    "id": "TC_SIM_08",
                    "desc": "Verify Longest Common Extension (LCE) greedy seed expansion",
                    "proc": "1: Create two documents with 45 identical words surrounded by dissimilar sentences.\n2: Run similarity engine.\n3: Inspect matching segment length.",
                    "out": "Initial 5-gram seed expands left and right to capture full 45-word identical passage.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "LCE greedy expansion successful"
                },
                {
                    "id": "TC_SIM_09",
                    "desc": "Verify contiguous span merging over minor 1-word insertion",
                    "proc": "1: Compare two texts identical except for single stop-word insertion.\n2: Run span merger with gap tolerance = 3 tokens.\n3: Inspect segment consolidation.",
                    "out": "Two adjacent matches merged into single composite highlighted block with gap recorded.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "Gap tolerance merging verified"
                },
                {
                    "id": "TC_SIM_10",
                    "desc": "Verify source and target document character offsets in matching_segments",
                    "proc": "1: Run compare_documents(p1, p2).\n2: Extract first segment: start_a, end_a, start_b, end_b, text.\n3: Verify slice equality: p1[start_a:end_a] == text.",
                    "out": "Slice p1[start_a:end_a] matches text and p2[start_b:end_b] matches text exactly.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "Slice coordinates verified"
                }
            ]
        },
        {
            "banner": "3. Exclusion Rules: Citations, Quotes & Small Match Thresholds",
            "cases": [
                {
                    "id": "TC_SIM_11",
                    "desc": "Verify exclusion of verbatim text enclosed in standard quotation marks",
                    "proc": "1: Include 40-word quote enclosed in double quotation marks (\"...\").\n2: Enable exclude_quotes = True.\n3: Check matching segment list.",
                    "out": "Quoted passage excluded from similarity tally; similarity score adjusted downward.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "Quotation exclusion verified"
                },
                {
                    "id": "TC_SIM_12",
                    "desc": "Verify bibliography and references section exclusion from peer matching",
                    "proc": "1: Compare two student papers sharing common course bibliography.\n2: Enable exclude_bibliography = True.\n3: Measure resulting similarity score.",
                    "out": "Common bibliographic citations excluded; only body essay content evaluated.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "Bibliography excluded cleanly"
                },
                {
                    "id": "TC_SIM_13",
                    "desc": "Verify small match threshold suppression (< 8 consecutive words)",
                    "proc": "1: Craft documents with accidental 4-word overlap ('in the context of').\n2: Run similarity comparison with min_match_words = 8.\n3: Check match count.",
                    "out": "Incidental 4-word overlap ignored; 0 false positive segments reported.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "Incidental matches filtered out"
                },
                {
                    "id": "TC_SIM_14",
                    "desc": "Verify asymmetric document comparison (Short excerpt copied into long paper)",
                    "proc": "1: Insert 100-word paragraph from Doc A (200 words total) into Doc B (2,000 words total).\n2: Compute directional similarity.\n3: Check scores.",
                    "out": "Doc A shows 50% overlap against Doc B, while Doc B shows 5% overlap against Doc A.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "Asymmetric ratios correctly reported"
                },
                {
                    "id": "TC_SIM_15",
                    "desc": "Verify collusion clustering across multiple student submissions in a team",
                    "proc": "1: Submit 4 papers where Student 2 and Student 3 share 65% content.\n2: Run cohort-level pairwise similarity matrix.\n3: Check collusion flags.",
                    "out": "Pairwise cluster clearly isolates Student 2 and Student 3 above 60% threshold.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "Collusion cluster correctly detected"
                }
            ]
        },
        {
            "banner": "4. Batch Processing Performance, Cross-Format & Edge Cases",
            "cases": [
                {
                    "id": "TC_SIM_16",
                    "desc": "Verify pairwise comparison execution time across 50 submissions (1,225 pairs)",
                    "proc": "1: Generate 50 student assignment submissions.\n2: Execute full pairwise batch comparison matrix.\n3: Measure total batch completion time.",
                    "out": "Entire pairwise matrix processes in < 8.0 seconds using inverted index pruning.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "Benchmark: 4.6s for 1,225 pairs"
                },
                {
                    "id": "TC_SIM_17",
                    "desc": "Verify cross-format similarity comparison (Digital PDF vs Scanned DOCX)",
                    "proc": "1: Submit identical text as native PDF and scanned image DOCX.\n2: Run OCR on DOCX and compare extracted text to PDF.\n3: Check similarity score.",
                    "out": "Similarity engine identifies match (> 90%) despite minor OCR transcription variance.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "Cross-format match > 92%"
                },
                {
                    "id": "TC_SIM_18",
                    "desc": "Verify mathematical equation and LaTeX symbol comparison",
                    "proc": "1: Compare two technical assignments containing LaTeX formulas ($E=mc^2$).\n2: Run tokenizer and similarity engine.\n3: Inspect formula segment matches.",
                    "out": "Mathematical token strings matched accurately between documents.",
                    "date": "12/10/2026",
                    "res": "Pass",
                    "note": "LaTeX tokens parsed cleanly"
                },
                {
                    "id": "TC_SIM_19",
                    "desc": "Verify side-by-side coordinate mapping for dual-document viewer UI",
                    "proc": "1: Fetch comparison report via /api/v1/analysis/{id}/comparison.\n2: Check coordinate payload for highlighted spans.\n3: Validate frontend mapping compatibility.",
                    "out": "JSON returns matching_segments with exact bounding boxes for synchronized dual-pane scrolling.",
                    "date": "12/10/2026",
                    "res": "Fail",
                    "note": "Defect VTX-312: High DPI scaling causes 4px coordinate offset in right viewer pane"
                },
                {
                    "id": "TC_SIM_20",
                    "desc": "Verify handling of multi-lingual submissions with mixed script (English + Latin/Devanagari)",
                    "proc": "1: Submit document containing mixed English and non-Latin character sets.\n2: Execute 5-gram tokenization and similarity comparison.\n3: Check tokenization stability.",
                    "out": "UTF-8 multi-byte characters tokenized cleanly without character clipping or out-of-range exceptions.",
                    "date": "12/10/2026",
                    "res": "Fail",
                    "note": "Defect VTX-329: Zero-width joiner characters cause irregular token window boundaries"
                }
            ]
        }
    ]

    # Create the 3 module sheets
    ws_m1 = create_module_test_sheet(
        sheet_title="CR101 - OCR Ingestion",
        module_code="CR101",
        module_name="Multi-Tier OCR Ingestion Engine",
        test_req="TR-OCR-01 - Adaptive Multi-Stage Document Ingestion & High-Fidelity Text Extraction",
        pass_count=19,
        fail_count=1,
        groups=ocr_groups
    )
    
    ws_m2 = create_module_test_sheet(
        sheet_title="CR102 - AI Detection",
        module_code="CR102",
        module_name="Turnitin-Aligned AI Writing Detection",
        test_req="TR-AID-02 - Explainable Multi-Factor AI Detection & Word-Weighted Proportionality",
        pass_count=19,
        fail_count=1,
        groups=aid_groups
    )
    
    ws_m3 = create_module_test_sheet(
        sheet_title="CR103 - Plagiarism Engine",
        module_code="CR103",
        module_name="Token-Indexed 5-Gram Plagiarism & Similarity",
        test_req="TR-SIM-03 - 5-Gram Token Inverted Indexing & Dual-Document Span Alignment",
        pass_count=18,
        fail_count=2,
        groups=sim_groups
    )
    
    # Auto-adjust column widths for all sheets
    all_sheets = [ws_cover, ws_report, ws_m1, ws_m2, ws_m3]
    for ws in all_sheets:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                # Ignore merged banner rows for width calculation
                if cell.row in [1, 2, 8, 9, 10] and ws.title in ["Record of Change", "Test Report"]:
                    continue
                val_str = str(cell.value or '')
                if '\n' in val_str:
                    lines = val_str.split('\n')
                    line_lens = [len(l) for l in lines]
                    max_len = max(max_len, max(line_lens))
                else:
                    max_len = max(max_len, len(val_str))
            
            # Custom column sizing based on sheet
            if ws.title in ["CR101 - OCR Ingestion", "CR102 - AI Detection", "CR103 - Plagiarism Engine"]:
                col_widths = {
                    "A": 14,  # ID
                    "B": 36,  # Description
                    "C": 48,  # Procedure
                    "D": 45,  # Expected Output
                    "E": 14,  # Date
                    "F": 12,  # Result
                    "G": 38   # Note
                }
                ws.column_dimensions[col_letter].width = col_widths.get(col_letter, 20)
            elif ws.title == "Record of Change":
                col_widths = {
                    "A": 18,
                    "B": 12,
                    "C": 26,
                    "D": 50,
                    "E": 20,
                    "F": 24,
                    "G": 22
                }
                ws.column_dimensions[col_letter].width = col_widths.get(col_letter, 20)
            elif ws.title == "Test Report":
                col_widths = {
                    "A": 8,
                    "B": 48,
                    "C": 12,
                    "D": 12,
                    "E": 14,
                    "F": 24
                }
                ws.column_dimensions[col_letter].width = col_widths.get(col_letter, 20)
                
    wb.save(filename)
    print(f"Workbook successfully created and saved to: {filename}")

if __name__ == "__main__":
    create_test_case_excel()
